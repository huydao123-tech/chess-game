"""
Script sinh file SRS.docx cho dự án Chess Game
Chạy: python generate_srs.py
"""
import subprocess, sys

# Auto-install python-docx if not present
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

from datetime import datetime
import os

OUTPUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SRS.docx")


def shade_table_header(table):
    for cell in table.rows[0].cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '1A3A5C')
        shd.set(qn('w:color'), 'FFFFFF')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True


def set_heading(doc, text, level=1, color=None):
    para = doc.add_paragraph(style=f"Heading {level}")
    run = para.add_run(text)
    run.font.bold = True
    if color:
        run.font.color.rgb = RGBColor(*color)
    return para


def add_table_row(table, cells, bold_first=False):
    row = table.add_row()
    for i, text in enumerate(cells):
        cell = row.cells[i]
        para = cell.paragraphs[0]
        para.clear()
        run = para.add_run(str(text))
        if bold_first and i == 0:
            run.bold = True
    return row


def create_srs():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # ── COVER PAGE ───────────────────────────────────────
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title_para.add_run("CHESS ONLINE")
    r.font.size = Pt(36); r.font.bold = True
    r.font.color.rgb = RGBColor(26, 58, 92)

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub_para.add_run("Software Requirements Specification (SRS)")
    r2.font.size = Pt(18); r2.font.color.rgb = RGBColor(70, 130, 180)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run(f"Version 1.0  |  {datetime.now().strftime('%d/%m/%Y')}")
    rm.font.size = Pt(12); rm.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()
    info_t = doc.add_table(rows=5, cols=2)
    info_t.style = 'Table Grid'
    info_rows = [
        ("Dự án:", "Chess Online - Web Application"),
        ("Ngôn ngữ:", "Java 21 (Spring Boot 3) / TypeScript (React 18 + Vite)"),
        ("Cơ sở dữ liệu:", "Microsoft SQL Server"),
        ("Tác giả:", "Chess Dev Team"),
        ("Cập nhật:", datetime.now().strftime('%d/%m/%Y %H:%M')),
    ]
    for i, (label, val) in enumerate(info_rows):
        info_t.rows[i].cells[0].text = label
        info_t.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        info_t.rows[i].cells[1].text = val

    doc.add_page_break()

    # ── 1. GIỚI THIỆU ─────────────────────────────────────
    set_heading(doc, "1. Giới thiệu", 1, (26, 58, 92))
    set_heading(doc, "1.1. Mục đích tài liệu", 2)
    doc.add_paragraph(
        "Tài liệu này mô tả Đặc tả Yêu cầu Phần mềm (SRS) cho hệ thống Chess Online – "
        "một ứng dụng web cho phép người dùng đăng nhập, quản lý hồ sơ và chơi cờ vua "
        "với Trí tuệ Nhân tạo (AI) ở nhiều cấp độ khó khác nhau. Tài liệu được tự động "
        "sinh và cập nhật bởi script generate_srs.py trong suốt vòng đời dự án."
    )

    set_heading(doc, "1.2. Phạm vi dự án", 2)
    doc.add_paragraph("Chess Online là nền tảng chơi cờ vua trực tuyến bao gồm:")
    scope_items = [
        "Hệ thống xác thực người dùng (đăng ký, đăng nhập, JWT token).",
        "Giao diện bàn cờ tương tác, tuân thủ đầy đủ luật cờ vua FIDE.",
        "Tích hợp engine AI Stockfish (Web Worker) với 20 cấp độ khó.",
        "Lưu trữ lịch sử ván đấu dạng PGN trong SQL Server.",
        "Hệ thống tính điểm ELO tự động sau mỗi ván đấu.",
        "Trang hồ sơ cá nhân với thống kê và lịch sử ván đấu.",
    ]
    for it in scope_items:
        doc.add_paragraph(it, style='List Bullet')

    set_heading(doc, "1.3. Định nghĩa & Thuật ngữ", 2)
    term_t = doc.add_table(rows=1, cols=2)
    term_t.style = 'Table Grid'
    term_t.rows[0].cells[0].text = "Thuật ngữ"
    term_t.rows[0].cells[1].text = "Định nghĩa"
    shade_table_header(term_t)
    terms = [
        ("FEN", "Forsyth–Edwards Notation – Chuỗi mô tả trạng thái bàn cờ."),
        ("PGN", "Portable Game Notation – Định dạng lưu trữ toàn bộ ván cờ."),
        ("Stockfish", "Engine cờ vua mã nguồn mở mạnh nhất thế giới."),
        ("ELO", "Hệ thống tính điểm rating cờ vua theo thuật toán Arpad Elo."),
        ("JWT", "JSON Web Token – Chuẩn xác thực stateless giữa client và server."),
        ("CORS", "Cross-Origin Resource Sharing – Cho phép frontend và backend khác port giao tiếp."),
        ("Web Worker", "Thread JavaScript chạy ngầm trong trình duyệt, dùng để chạy Stockfish."),
    ]
    for t, d in terms:
        add_table_row(term_t, [t, d], bold_first=True)
    doc.add_paragraph()

    # ── 2. KIẾN TRÚC ──────────────────────────────────────
    set_heading(doc, "2. Kiến trúc Hệ thống", 1, (26, 58, 92))
    doc.add_paragraph(
        "Hệ thống được xây dựng theo mô hình 3 tầng (Three-Tier Architecture): "
        "Presentation Layer (React), Business Logic Layer (Spring Boot), Data Layer (SQL Server)."
    )
    set_heading(doc, "2.1. Sơ đồ Kiến trúc", 2)
    arch_p = doc.add_paragraph()
    arch_r = arch_p.add_run(
        "[ Browser ]\n"
        "  React 18 + TypeScript (Vite)\n"
        "  chess.js  |  react-chessboard  |  stockfish.js (Web Worker)\n"
        "      |\n"
        "      |  REST API (HTTP/JSON + JWT)\n"
        "      v\n"
        "[ Spring Boot 3 - Port 8080 ]\n"
        "  Spring Web  |  Spring Security  |  Spring Data JPA\n"
        "      |\n"
        "      |  JDBC (Microsoft JDBC Driver)\n"
        "      v\n"
        "[ SQL Server Database ]\n"
        "  Tables: users | games"
    )
    arch_r.font.name = 'Courier New'
    arch_r.font.size = Pt(9)
    arch_p.paragraph_format.left_indent = Cm(1)

    set_heading(doc, "2.2. Cấu trúc Thư mục Dự án", 2)
    dir_p = doc.add_paragraph()
    dir_r = dir_p.add_run(
        "ChessGame/\n"
        "  SRS.docx                     <- Tài liệu này\n"
        "  generate_srs.py              <- Script tự động tạo SRS\n"
        "  chess-backend/               <- Spring Boot\n"
        "    pom.xml\n"
        "    src/main/java/com/chess/\n"
        "      config/                  (CORS, Security, JWT)\n"
        "      controllers/             (AuthController, GameController, UserController)\n"
        "      models/                  (User, Game entities)\n"
        "      repositories/\n"
        "      services/\n"
        "      dto/                     (Request/Response DTOs)\n"
        "    src/main/resources/\n"
        "      application.properties\n"
        "  chess-frontend/              <- React + Vite\n"
        "    src/\n"
        "      api/                     (axios instances + interceptors)\n"
        "      components/              (ChessBoard, Navbar, Timer, MoveList...)\n"
        "      pages/                   (Home, Login, Register, Game, Profile, History)\n"
        "      store/                   (Zustand auth + game state)\n"
        "      workers/                 (stockfish.worker.ts)\n"
        "      styles/                  (global.css + CSS modules)\n"
        "    public/stockfish/          (stockfish.js engine files)"
    )
    dir_r.font.name = 'Courier New'
    dir_r.font.size = Pt(9)
    dir_p.paragraph_format.left_indent = Cm(1)
    doc.add_paragraph()

    # ── 3. YÊU CẦU CHỨC NĂNG ─────────────────────────────
    set_heading(doc, "3. Yêu cầu Chức năng", 1, (26, 58, 92))
    fr_t = doc.add_table(rows=1, cols=4)
    fr_t.style = 'Table Grid'
    for i, h in enumerate(["ID", "Tính năng", "Mô tả", "Ưu tiên"]):
        fr_t.rows[0].cells[i].text = h
    shade_table_header(fr_t)
    reqs = [
        ("FR-01", "Đăng ký", "Nhập email, username, password. Kiểm tra trùng, mã hóa BCrypt, lưu DB.", "Cao"),
        ("FR-02", "Đăng nhập", "Xác thực username/password, trả về JWT Access+Refresh Token.", "Cao"),
        ("FR-03", "Refresh Token", "Tự động gia hạn Access Token qua Refresh Token.", "Cao"),
        ("FR-04", "Đăng xuất", "Xóa token client, vô hiệu hoá Refresh Token server.", "Cao"),
        ("FR-05", "Xem hồ sơ", "Hiển thị username, email, ELO, thống kê thắng/thua/hòa.", "Trung bình"),
        ("FR-06", "Chỉnh sửa hồ sơ", "Cập nhật username, avatar.", "Thấp"),
        ("FR-07", "Chọn cấp độ AI", "1 (Beginner) đến 20 (Grandmaster - Stockfish max depth).", "Cao"),
        ("FR-08", "Chọn màu quân", "Chơi quân Trắng, Đen hoặc Ngẫu nhiên.", "Cao"),
        ("FR-09", "Chơi cờ vs AI", "Bàn cờ kéo-thả, highlight nước hợp lệ, AI phản hồi qua Stockfish.", "Cao"),
        ("FR-10", "Đồng hồ bấm giờ", "Chọn 1/3/5/10/15/30 phút mỗi bên. Hết giờ thua cuộc.", "Cao"),
        ("FR-11", "Kết thúc ván", "Phát hiện Checkmate, Stalemate, Draw (50 nước, 3 lần lặp). Dialog kết quả.", "Cao"),
        ("FR-12", "Lưu ván đấu", "Sau ván kết thúc, lưu PGN, kết quả, ELO change vào DB.", "Cao"),
        ("FR-13", "Lịch sử ván đấu", "Danh sách ván đấu có phân trang, filter theo kết quả.", "Trung bình"),
        ("FR-14", "Replay ván cờ", "Xem lại ván đã lưu từng nước đi (Prev/Next/Auto).", "Trung bình"),
        ("FR-15", "Cập nhật ELO", "Tính ELO theo công thức chuẩn sau mỗi ván với AI.", "Trung bình"),
    ]
    for r in reqs:
        add_table_row(fr_t, list(r))
    doc.add_paragraph()

    # ── 4. YÊU CẦU PHI CHỨC NĂNG ─────────────────────────
    set_heading(doc, "4. Yêu cầu Phi chức năng", 1, (26, 58, 92))
    nfr_t = doc.add_table(rows=1, cols=3)
    nfr_t.style = 'Table Grid'
    for i, h in enumerate(["Loại", "Yêu cầu", "Chỉ tiêu"]):
        nfr_t.rows[0].cells[i].text = h
    shade_table_header(nfr_t)
    nfrs = [
        ("Hiệu năng", "Thời gian phản hồi API", "< 300ms cho 95% requests"),
        ("Hiệu năng", "AI phản hồi nước đi", "< 2 giây (chạy Web Worker không block UI)"),
        ("Hiệu năng", "Render bàn cờ", "60 FPS, không giật lag"),
        ("Bảo mật", "Xác thực JWT", "HS256, Access Token 15 phút, BCrypt strength 12"),
        ("Bảo mật", "Input Validation", "Validate Backend, JPA chống SQL Injection"),
        ("Khả năng mở rộng", "Stateless Backend", "Không lưu session server, dễ scale"),
        ("Tương thích", "Trình duyệt", "Chrome 90+, Firefox 88+, Edge 90+"),
        ("Tương thích", "Responsive", "Tối ưu Desktop 1280px+, hỗ trợ tablet"),
    ]
    for r in nfrs:
        add_table_row(nfr_t, list(r))
    doc.add_paragraph()

    # ── 5. DATABASE SCHEMA ────────────────────────────────
    set_heading(doc, "5. Thiết kế Cơ sở Dữ liệu", 1, (26, 58, 92))

    set_heading(doc, "5.1. Bảng users", 2)
    u_t = doc.add_table(rows=1, cols=4)
    u_t.style = 'Table Grid'
    for i, h in enumerate(["Cột", "Kiểu dữ liệu", "Ràng buộc", "Mô tả"]):
        u_t.rows[0].cells[i].text = h
    shade_table_header(u_t)
    for r in [
        ("id", "BIGINT", "PK, IDENTITY", "Khóa chính"),
        ("username", "NVARCHAR(50)", "UNIQUE, NOT NULL", "Tên đăng nhập"),
        ("email", "NVARCHAR(100)", "UNIQUE, NOT NULL", "Email"),
        ("password_hash", "NVARCHAR(255)", "NOT NULL", "BCrypt hash"),
        ("elo_rating", "INT", "DEFAULT 1200", "Điểm ELO"),
        ("total_games", "INT", "DEFAULT 0", "Tổng ván"),
        ("wins", "INT", "DEFAULT 0", "Số thắng"),
        ("losses", "INT", "DEFAULT 0", "Số thua"),
        ("draws", "INT", "DEFAULT 0", "Số hòa"),
        ("avatar_url", "NVARCHAR(255)", "NULL", "URL avatar"),
        ("created_at", "DATETIME2", "DEFAULT GETDATE()", "Ngày tạo"),
        ("refresh_token", "NVARCHAR(512)", "NULL", "JWT Refresh Token"),
        ("is_active", "BIT", "DEFAULT 1", "Tài khoản active?"),
    ]:
        add_table_row(u_t, list(r), bold_first=True)
    doc.add_paragraph()

    set_heading(doc, "5.2. Bảng games", 2)
    g_t = doc.add_table(rows=1, cols=4)
    g_t.style = 'Table Grid'
    for i, h in enumerate(["Cột", "Kiểu dữ liệu", "Ràng buộc", "Mô tả"]):
        g_t.rows[0].cells[i].text = h
    shade_table_header(g_t)
    for r in [
        ("id", "BIGINT", "PK, IDENTITY", "Khóa chính"),
        ("user_id", "BIGINT", "FK → users(id)", "ID người chơi"),
        ("player_color", "NVARCHAR(5)", "NOT NULL", "WHITE hoặc BLACK"),
        ("ai_level", "INT", "NOT NULL (1-20)", "Cấp độ Stockfish"),
        ("result", "NVARCHAR(10)", "NOT NULL", "WIN / LOSS / DRAW"),
        ("pgn", "NVARCHAR(MAX)", "NULL", "Toàn bộ ván cờ PGN"),
        ("final_fen", "NVARCHAR(100)", "NULL", "FEN trạng thái cuối"),
        ("total_moves", "INT", "DEFAULT 0", "Tổng số nước đi"),
        ("duration_seconds", "INT", "NULL", "Thời gian ván (giây)"),
        ("time_control", "INT", "NULL", "Thời gian kiểm soát (giây)"),
        ("elo_before", "INT", "NULL", "ELO trước ván"),
        ("elo_after", "INT", "NULL", "ELO sau ván"),
        ("elo_change", "INT", "NULL", "Thay đổi ELO +/-"),
        ("played_at", "DATETIME2", "DEFAULT GETDATE()", "Thời điểm chơi"),
    ]:
        add_table_row(g_t, list(r), bold_first=True)
    doc.add_paragraph()

    # ── 6. REST API ───────────────────────────────────────
    set_heading(doc, "6. Đặc tả REST API", 1, (26, 58, 92))
    api_t = doc.add_table(rows=1, cols=5)
    api_t.style = 'Table Grid'
    for i, h in enumerate(["Method", "Endpoint", "Auth", "Mô tả", "Response"]):
        api_t.rows[0].cells[i].text = h
    shade_table_header(api_t)
    for r in [
        ("POST", "/api/auth/register", "None", "Đăng ký tài khoản", "201 + UserDTO"),
        ("POST", "/api/auth/login", "None", "Đăng nhập, nhận JWT", "200 + TokenDTO"),
        ("POST", "/api/auth/refresh", "None", "Làm mới Access Token", "200 + TokenDTO"),
        ("POST", "/api/auth/logout", "JWT", "Đăng xuất", "204"),
        ("GET", "/api/users/me", "JWT", "Lấy profile hiện tại", "200 + UserDTO"),
        ("PUT", "/api/users/me", "JWT", "Cập nhật profile", "200 + UserDTO"),
        ("GET", "/api/users/me/stats", "JWT", "Thống kê ELO, win rate", "200 + StatsDTO"),
        ("GET", "/api/games", "JWT", "Lịch sử ván đấu (paged)", "200 + Page<GameDTO>"),
        ("POST", "/api/games", "JWT", "Lưu kết quả ván mới", "201 + GameDTO"),
        ("GET", "/api/games/{id}", "JWT", "Chi tiết ván + PGN", "200 + GameDetailDTO"),
    ]:
        add_table_row(api_t, list(r))
    doc.add_paragraph()

    # ── 7. UI SPECIFICATION ───────────────────────────────
    set_heading(doc, "7. Đặc tả Giao diện Người dùng", 1, (26, 58, 92))
    set_heading(doc, "7.1. Design System", 2)
    doc.add_paragraph("Ứng dụng dùng thiết kế Dark Mode hiện đại:")
    for it in [
        "Màu nền: #0A0E1A (Deep Navy Black)",
        "Màu nhấn: #5B8AF0 (Royal Blue)",
        "Màu thành công: #4CAF82 (Emerald Green)",
        "Font: 'Inter' từ Google Fonts",
        "Hiệu ứng: Glassmorphism panels, subtle glow, smooth animations",
    ]:
        doc.add_paragraph(it, style='List Bullet')

    set_heading(doc, "7.2. Danh sách Trang", 2)
    pg_t = doc.add_table(rows=1, cols=3)
    pg_t.style = 'Table Grid'
    for i, h in enumerate(["Trang", "Route", "Mô tả"]):
        pg_t.rows[0].cells[i].text = h
    shade_table_header(pg_t)
    for r in [
        ("Landing Page", "/", "Trang giới thiệu với hero section, CTA button."),
        ("Đăng nhập", "/login", "Form đăng nhập, validation realtime."),
        ("Đăng ký", "/register", "Form đăng ký, password strength indicator."),
        ("Chơi vs AI", "/play", "Bàn cờ, AI settings panel, timer, move history."),
        ("Lịch sử", "/history", "Bảng ván đấu, filter, ELO chart."),
        ("Hồ sơ", "/profile", "Thông tin cá nhân, thống kê."),
        ("Replay", "/replay/:id", "Xem lại ván cờ từng bước."),
    ]:
        add_table_row(pg_t, list(r))
    doc.add_paragraph()

    # ── 8. KẾ HOẠCH ──────────────────────────────────────
    set_heading(doc, "8. Kế hoạch Triển khai", 1, (26, 58, 92))
    ph_t = doc.add_table(rows=1, cols=4)
    ph_t.style = 'Table Grid'
    for i, h in enumerate(["Giai đoạn", "Nội dung", "Ưu tiên", "Trạng thái"]):
        ph_t.rows[0].cells[i].text = h
    shade_table_header(ph_t)
    for r in [
        ("Phase 0", "Khởi tạo project, SRS.docx", "Cao", "Hoàn thành"),
        ("Phase 1", "Backend: Auth API + JWT + SQL Server", "Cao", "Dang lam"),
        ("Phase 2", "Backend: Game API + ELO calculation", "Cao", "Cho"),
        ("Phase 3", "Frontend: Landing, Login, Register", "Cao", "Cho"),
        ("Phase 4", "Frontend: Game page + Stockfish AI", "Cao", "Cho"),
        ("Phase 5", "Frontend: History, Profile, Replay", "TB", "Cho"),
        ("Phase 6", "Tich hop Frontend Backend, test", "Cao", "Cho"),
        ("Phase 7", "Toi uu, polish UI/UX", "TB", "Cho"),
    ]:
        add_table_row(ph_t, list(r))
    doc.add_paragraph()

    # ── 9. VERSION HISTORY ────────────────────────────────
    set_heading(doc, "9. Lịch sử Phiên bản", 1, (26, 58, 92))
    ver_t = doc.add_table(rows=1, cols=4)
    ver_t.style = 'Table Grid'
    for i, h in enumerate(["Phiên bản", "Ngày", "Tác giả", "Thay đổi"]):
        ver_t.rows[0].cells[i].text = h
    shade_table_header(ver_t)
    add_table_row(ver_t, ["1.0", datetime.now().strftime('%d/%m/%Y'), "Chess Dev Team",
                          "Phiên bản đầu tiên - Khởi tạo SRS đầy đủ."])

    doc.save(OUTPUT_PATH)
    print(f"SRS.docx created at: {OUTPUT_PATH}")


if __name__ == "__main__":
    create_srs()
